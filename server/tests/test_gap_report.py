from datetime import date
from types import SimpleNamespace
from uuid import UUID, uuid4

import pytest
import sqlalchemy as sa
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.config.database import get_db
from app.models.entities import Eco
from app.services.gap_report_service import GapReportService
from app.utils.security import current_user
from main import app


START = date(2026, 10, 1)
END = date(2026, 10, 10)
URL = "/api/v1/reports/gap"
PARAMS = {"start_date": START.isoformat(), "end_date": END.isoformat()}


@pytest.fixture
def db():
    engine = sa.create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    Eco.__table__.create(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()


def add_eco(db, **overrides):
    values = dict(
        id=uuid4(), eco="ECO-TEST", owner="Owner", status="RELEASED",
        eco_type="STANDARD",
        az_eco_creation_date=date(2026, 8, 15),
        effective_date=date(2026, 9, 20),
        second_aprov_rd_finish_1=date(2026, 10, 5), gap_14d_note="Specific note",
    )
    values.update(overrides)
    row = Eco(**values)
    db.add(row)
    db.commit()
    return row


@pytest.mark.parametrize(("overrides", "included", "gap", "requires"), [
    ({}, True, 15, True),
    ({"second_aprov_rd_finish_1": START}, True, 11, False),
    ({"second_aprov_rd_finish_1": END}, True, 20, True),
    ({"second_aprov_rd_finish_1": date(2026, 9, 30)}, False, None, False),
    ({"second_aprov_rd_finish_1": date(2026, 10, 11)}, False, None, False),
    ({"status": "WORKING"}, False, None, False),
    ({"status": " released "}, True, 15, True),
    ({"status": "\tReLeAsEd\n"}, True, 15, True),
    ({"status": None}, False, None, False),
    ({"second_aprov_rd_finish_1": None, "second_aprov_rd_finish_2": START}, False, None, False),
    ({"second_aprov_rd_finish_2": date(2026, 11, 1)}, True, 15, True),
    ({"effective_date": None}, True, None, False),
    ({"effective_date": date(2026, 9, 21)}, True, 14, False),
    ({"effective_date": date(2026, 10, 6)}, True, -1, False),
])
def test_report_rules(db, overrides, included, gap, requires):
    row = add_eco(db, **overrides)
    result = GapReportService(db).report(START, END)
    assert result.start_date == START
    assert result.end_date == END
    assert result.total == len(result.items) == int(included)
    if included:
        item = result.items[0]
        assert item.id == str(row.id)
        assert item.eco_type == row.eco_type
        assert item.approval_date == row.second_aprov_rd_finish_1
        assert item.az_eco_creation_date == row.az_eco_creation_date
        assert item.effective_date == row.effective_date
        assert item.gap_14d == gap
        assert item.requires_gap_14d_note is requires
        assert item.gap_14d_note == "Specific note"


def test_report_order_is_deterministic(db):
    for code, approval, ident in [
        (None, END, 5), ("B", END, 4), ("A", END, 3),
        ("Z", START, 1), ("A", END, 2),
    ]:
        add_eco(db, eco=code, second_aprov_rd_finish_1=approval, id=UUID(int=(10 << 124) + ident))
    result = GapReportService(db).report(START, END)
    assert [item.id for item in result.items] == [str(UUID(int=(10 << 124) + n)) for n in [1, 2, 3, 4, 5]]


def test_report_rejects_inverted_period_before_query():
    with pytest.raises(HTTPException) as exc:
        GapReportService(None).report(END, START)
    assert exc.value.status_code == 422
    assert "Data inicial" in exc.value.detail


def test_single_day_and_empty_report(db):
    assert GapReportService(db).report(START, END).total == 0
    add_eco(db, second_aprov_rd_finish_1=START)
    assert GapReportService(db).report(START, START).total == 1


@pytest.fixture
def client(db):
    before = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[current_user] = lambda: SimpleNamespace(role="analyst")
    try:
        with TestClient(app) as client:
            yield client
    finally:
        app.dependency_overrides = before


def test_report_http_response(client, db):
    row = add_eco(db)
    response = client.get(URL, params=PARAMS)
    assert response.status_code == 200
    assert response.json() == {
        "start_date": "2026-10-01", "end_date": "2026-10-10", "total": 1,
        "items": [{
            "id": str(row.id), "eco": "ECO-TEST", "eco_type": "STANDARD",
            "az_eco_creation_date": "2026-08-15",
            "effective_date": "2026-09-20", "approval_date": "2026-10-05",
            "gap_14d": 15, "gap_14d_note": "Specific note",
            "requires_gap_14d_note": True,
        }],
    }


@pytest.mark.parametrize("params", [
    {}, {"start_date": "2026-10-01"}, {"end_date": "2026-10-10"},
    {"start_date": "invalid", "end_date": "2026-10-10"},
    {"start_date": "2026-10-01", "end_date": "2026-02-30"},
    {"start_date": "2026-10-10", "end_date": "2026-10-01"},
])
def test_report_http_invalid_params(client, params):
    assert client.get(URL, params=params).status_code == 422


def test_report_http_null_effective_date(client, db):
    add_eco(db, az_eco_creation_date=None, effective_date=None, gap_14d_note=None)
    response = client.get(URL, params=PARAMS)
    assert response.status_code == 200
    item = response.json()["items"][0]
    assert item["az_eco_creation_date"] is None
    assert item["effective_date"] is None
    assert item["gap_14d"] is None
    assert item["gap_14d_note"] is None
    assert item["requires_gap_14d_note"] is False


def test_report_requires_authentication(db):
    before = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides.pop(current_user, None)
    try:
        response = TestClient(app).get(URL, params=PARAMS)
    finally:
        app.dependency_overrides = before
    assert response.status_code == 403  # Existing HTTPBearer behavior without token.


def test_report_rejects_invalid_token(db):
    before = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides.pop(current_user, None)
    try:
        response = TestClient(app).get(URL, params=PARAMS, headers={"Authorization": "Bearer invalid"})
    finally:
        app.dependency_overrides = before
    assert response.status_code == 401


@pytest.mark.parametrize("path", [URL, URL + "/export"])
def test_export_auth_required(db, path):
    before = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides.pop(current_user, None)
    try:
        response = TestClient(app).get(path, params=PARAMS)
    finally:
        app.dependency_overrides = before
    assert response.status_code == 403


@pytest.mark.parametrize("params", [{}, {"start_date": "2026-10-10", "end_date": "2026-10-01"}])
def test_export_invalid_params(client, params):
    assert client.get(URL + "/export", params=params).status_code == 422


def test_gap_export_valid_workbook_matches_report(client, db):
    from io import BytesIO
    from openpyxl import load_workbook

    add_eco(db, eco="START", second_aprov_rd_finish_1=START)
    add_eco(db, eco="END", second_aprov_rd_finish_1=END, az_eco_creation_date=None, effective_date=None, gap_14d_note=None)
    add_eco(db, eco="NEGATIVE", effective_date=date(2026, 10, 6))
    add_eco(db, eco="NORMALIZED", status=" released ")
    add_eco(db, eco="BEFORE", second_aprov_rd_finish_1=date(2026, 9, 30))
    add_eco(db, eco="AFTER", second_aprov_rd_finish_1=date(2026, 10, 11))
    add_eco(db, eco="WORKING", status="WORKING")
    add_eco(db, eco="FINISH2", second_aprov_rd_finish_1=None, second_aprov_rd_finish_2=START)
    add_eco(db, eco="=1+1", gap_14d_note="=HYPERLINK(123)")
    report = client.get(URL, params=PARAMS).json()
    response = client.get(URL + "/export", params=PARAMS)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"] == 'attachment; filename="relatorio_gap_2026-10-01_2026-10-10.xlsx"'
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    headers = [c.value for c in sheet[1]]
    assert headers == ["ECO", "ECO TYPE", "Effective Date", "AZ ECO CREATION DATE", "Approval Date", "LEAD TIME", "Coment\u00e1rio / Justificativa"]
    assert "GAP 14D" not in headers
    assert sheet.max_row == report["total"] + 1
    for cells, item in zip(sheet.iter_rows(min_row=2), report["items"]):
        assert cells[0].value == item["eco"]
        assert cells[1].value == item["eco_type"]
        assert (cells[2].value.date().isoformat() if cells[2].value else None) == item["effective_date"]
        assert (cells[3].value.date().isoformat() if cells[3].value else None) == item["az_eco_creation_date"]
        assert cells[4].value.date().isoformat() == item["approval_date"]
        assert all(cells[index].number_format == "dd/mm/yyyy" for index in (2, 3, 4) if cells[index].value)
        assert cells[5].value == item["gap_14d"]
        assert cells[5].data_type == "n"
        assert cells[6].value == item["gap_14d_note"]
        assert cells[0].data_type != "f"
        assert cells[6].data_type != "f"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:G{sheet.max_row}"
    workbook.close()
    assert client.get(URL, params=PARAMS).json() == report


def test_gap_export_empty_workbook(client):
    from io import BytesIO
    from openpyxl import load_workbook
    response = client.get(URL + "/export", params=PARAMS)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.active.max_row == 1
    assert workbook.active.max_column == 7
    workbook.close()
