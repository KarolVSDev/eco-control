import io

from datetime import date
from types import SimpleNamespace

import pytest

from fastapi import HTTPException

from openpyxl import Workbook, load_workbook

from app.services.eco_import_service import (
    EcoImportService,
)


# =========================================================
# CABEÇALHOS REAIS DA CTRL GERAL
# =========================================================

CTRL_GERAL_HEADERS = {
    1: "GAP",
    2: "ITEM",
    3: "MONTH",
    4: "PRODUCT",
    5: "AU",
    6: "OBU",
    7: "GROUP",
    8: "OWNER",
    9: "ITEM TYPE",
    10: "ECO TYPE",
    11: "Change BOM?",
    12: "STATUS",
    13: "RECEB.",
    14: "ECO",
    15: "CHANGE REASON",
    16: "HQ ECO RELEASE DATE",
    17: "AZ ECO REGISTER DATE",
    18: "Auto ECR?",
    19: "1-Delay ECO Register",
    20: "ECO registration week",
    21: "Council meeting Week",

    22: "HZ/IN/KR\nECO(1)",
    23: "HZ/IN/KR\nRECEIVE DATE (1)",
    24: "ECO Origem (1)",

    25: "HZ/IN/KR\nECO(2)",
    26: "HZ/SH/IN\nRECEIVE DATE (2)",
    27: "ECO Origem (2)",

    28: "HZ/IN/KR\nECO(3)",
    29: "HZ/SH/IN\nRECEIVE DATE (3)",
    30: "ECO Origem (3)",

    31: "AZ_ECO_NO.",
    32: "Change reason2",
    33: "Change BOM?",
    34: "AZ ECO CREATION DATE",
    35: "GAP Start AZ ECO",
    36: "Contar Eco emitida > 1 dias",
    37: "Effective Date",
    38: "GAP AGREEMENT",

    39: "AGREEMENT\nOTHER\nDEPTS START\n(1)",
    40: "AGREEMENT\nOTHER\nDEPTS FINSHI\n(1)",
    41: "GAP AGREEMENT\nOTHER\nDEPTS (1)",

    42: "AGREEMENT\nOTHER\nDEPTS START\n(2)",
    43: "AGREEMENT\nOTHER\nDEPTS FINSHI\n(2)",
    44: "GAP AGREEMENT\nOTHER\nDEPTS (2)",

    45: "AGREEMENT\nOTHER\nDEPTS START\n(3)",
    46: "AGREEMENT\nOTHER\nDEPTS FINSHI\n(3)",
    47: "GAP AGREEMENT\nOTHER\nDEPTS (3)",

    48: "GAP TOTAL AGREEMENT\nOTHER\nDEPTS",

    49: "2ST APROV R&D",
    50: "2ST APROV R&D START (1)",
    51: "2ST APROV R&D FINISH (1)",
    52: "GAP 2ST APROV R&D (1)",

    53: "2ST APROV R&D START (2)",
    54: "2ST APROV R&D FINISH (2)",
    55: "GAP 2ST APROV R&D (2)",
    56: "GAP TOTAL 2ST APROV R&D",

    57: "ECO release week",
    58: "AZ Gap",
    59: "Contar Eco concluída > 7 dias",
    60: "Contar Eco concluída > 10 dias",
    61: "GAP 14D",
    62: "Total Gap",
    63: "RELEASE MONTH",
    64: "RELEASE YEAR",

    65: "CHANGE REASON",
    66: "MODEL",
    67: "New Model?\n(Check in NPI plan 26Y)",
    68: (
        "Origem approval\n"
        "(mandatory for local and AV ECO)"
    ),
    69: "Event",
    70: "COMMENTS",

    71: "SET ECO",
    72: "SET ECO REGISTER DATE",
    73: "SET ECO RELEASE DATE",
    74: "CHANGE REASON",
    75: "MODEL",
}


# =========================================================
# HELPERS
# =========================================================

def create_service(
    existing=None,
    obu_map=None,
    owner_map=None,
):
    user = SimpleNamespace(
        role="admin",
        email="admin@eco.com",
        full_name="Administrador",
    )


    service = EcoImportService(
        db=None,
        user=user,
    )


    service._settings_maps = lambda: (
        obu_map
        or {
            "NW1": "GLZ",
        },

        owner_map
        or {
            "owner.test": "BOM",
        },
    )


    service._existing_rows_by_eco = (
        lambda eco_codes:
            existing
            or {}
    )


    return service


def create_workbook(
    *,
    sheet_name="CTRL GERAL",
    rows=None,
):
    workbook = Workbook()


    worksheet = (
        workbook.active
    )


    worksheet.title = (
        sheet_name
    )


    # -----------------------------------------------------
    # CABEÇALHO REAL NA LINHA 7
    # -----------------------------------------------------

    header_row = 7


    for (
        column,
        value,
    ) in CTRL_GERAL_HEADERS.items():

        worksheet.cell(
            row=header_row,
            column=column,
            value=value,
        )


    # -----------------------------------------------------
    # DADOS
    # -----------------------------------------------------

    if rows is None:

        rows = [
            {
                "item": 1,

                "product": "AAA",

                "obu": "NW1",

                "owner":
                    "owner.test",

                "item_type":
                    "MEC",

                "eco_type":
                    "REGULAR",

                "change_bom":
                    "YES",

                "status":
                    "RELEASED",

                "receb":
                    "NORMAL",

                "eco":
                    "ECO-001",

                "change_reason":
                    "Alteração teste",

                "hq_eco_release_date":
                    date(
                        2026,
                        1,
                        2,
                    ),

                "az_eco_register_date":
                    date(
                        2026,
                        1,
                        5,
                    ),

                "auto_ecr":
                    "NO",

                "az_eco_creation_date":
                    date(
                        2026,
                        1,
                        6,
                    ),

                "second_aprov_rd":
                    "R&D",

                "second_aprov_rd_start_1":
                    date(
                        2026,
                        1,
                        7,
                    ),

                "second_aprov_rd_finish_1":
                    date(
                        2026,
                        1,
                        12,
                    ),

                "effective_date": date(2026, 2, 4),

                "comments":
                    "Teste de importação",
            }
        ]


    field_to_column = {
        spec["field"]: column

        for (
            column,
            spec,
        ) in (
            EcoImportService
            .COLUMN_SPECS
            .items()
        )
    }


    current_row = 8


    for data in rows:

        worksheet.cell(
            row=current_row,
            column=2,
            value=data.get(
                "item"
            ),
        )


        for (
            field,
            value,
        ) in data.items():

            if field == "item":
                continue


            column = (
                field_to_column.get(
                    field
                )
            )


            if column is None:
                continue


            worksheet.cell(
                row=current_row,
                column=column,
                value=value,
            )


        current_row += 1


    output = io.BytesIO()


    workbook.save(
        output
    )


    workbook.close()


    return output.getvalue()


# =========================================================
# ESTRUTURA DAS COLUNAS
# =========================================================

def test_rd_columns_match_ctrl_geral():

    specs = (
        EcoImportService
        .COLUMN_SPECS
    )


    assert (
        specs[49]["field"]
        ==
        "second_aprov_rd"
    )


    assert (
        specs[50]["field"]
        ==
        "second_aprov_rd_start_1"
    )


    assert (
        specs[51]["field"]
        ==
        "second_aprov_rd_finish_1"
    )


    assert (
        specs[53]["field"]
        ==
        "second_aprov_rd_start_2"
    )


    assert (
        specs[54]["field"]
        ==
        "second_aprov_rd_finish_2"
    )


# =========================================================
# PREVIEW DE UMA ECO NOVA
# =========================================================

def test_preview_new_eco():

    service = (
        create_service()
    )


    file_bytes = (
        create_workbook()
    )


    result = (
        service.preview(
            file_bytes,
            "controle.xlsx",
        )
    )


    assert (
        result["sheet"]
        ==
        "CTRL GERAL"
    )


    assert (
        result["header_row"]
        == 7
    )


    assert (
        result["total_rows"]
        == 1
    )


    assert (
        result["valid_rows"]
        == 1
    )


    assert (
        result["new_rows"]
        == 1
    )


    assert (
        result["update_rows"]
        == 0
    )


    assert (
        result["error_rows"]
        == 0
    )


    assert result["preview_mode"] == "ALL"
    assert result["preview_rows_total"] == 1
    assert result["preview_rows_count"] == 1
    assert result["preview_truncated"] is False


    assert (
        result["can_import"]
        is True
    )


    row = (
        result[
            "rows"
        ][0]
    )


    assert (
        row["action"]
        ==
        "NEW"
    )


    assert (
        row["source_item"]
        == 1
    )


    assert (
        row["eco"]
        ==
        "ECO-001"
    )


    assert (
        row["data"]["obu"]
        ==
        "NW1"
    )


    assert (
        row["derived"]["au"]
        ==
        "GLZ"
    )


    assert (
        row["derived"]["group"]
        ==
        "BOM"
    )


    assert (
        row["derived"]["month"]
        ==
        "JAN"
    )


def test_preview_uses_internal_month_and_sheet_release_month():

    rows = [
        {
            "item": 1,
            "obu": "NW1",
            "owner": "owner.test",
            "item_type": "MEC",
            "eco_type": "REGULAR",
            "status": "WORKING",
            "eco": "ECO-FEV",
            "second_aprov_rd_finish_1": date(2026, 2, 7),
            "az_eco_register_date": date(2026, 2, 5),
        }
    ]

    result = create_service().preview(
        create_workbook(rows=rows),
        "controle.xlsx",
    )
    derived = result["rows"][0]["derived"]

    assert derived["month"] == "FEB"
    assert derived["release_month"] == "FEV"


def test_record_import_history_serializes_populated_fields():

    service = create_service()
    history_records = []
    service.history.add = history_records.append
    eco = SimpleNamespace(
        id="eco-id",
        eco="ECO-HISTORY",
        item=7,
        position=3,
        month="FEB",
        au="GLZ",
        group_name="BOM",
    )

    service._record_import_history(
        eco,
        {
            "product": "AAA",
            "az_eco_register_date": date(2026, 2, 10),
            "empty_value": None,
        },
    )

    records_by_field = {
        record.field_key: record
        for record in history_records
    }

    assert records_by_field["az_eco_register_date"].new_value == "2026-02-10"
    assert records_by_field["group_name"].field_label == "GROUP"
    assert records_by_field["product"].new_value == "AAA"
    assert "empty_value" not in records_by_field
    assert all(record.action == "created" for record in history_records)


def test_import_file_persists_only_new_rows_and_commits_once():

    existing_row = {
        "item": 1,
        "product": "AAA",
        "obu": "NW1",
        "owner": "owner.test",
        "item_type": "MEC",
        "eco_type": "REGULAR",
        "change_bom": "YES",
        "status": "WORKING",
        "receb": "NORMAL",
        "eco": "ECO-IMPORT",
        "change_reason": "Alteração teste",
        "hq_eco_release_date": date(2026, 1, 2),
        "az_eco_register_date": date(2026, 1, 5),
    }
    baseline = create_service().preview(
        create_workbook(rows=[existing_row]),
        "controle.xlsx",
    )
    existing_record = SimpleNamespace(
        **baseline["rows"][0]["data"],
        id="existing-id",
        item=1,
    )
    new_row = {
        **existing_row,
        "item": 2,
        "product": "BBB",
    }
    service = create_service(
        existing={
            "ECO-IMPORT": [existing_record],
        }
    )
    persisted = []
    histories = []
    commits = []
    rollbacks = []
    service.db = SimpleNamespace(
        add=persisted.append,
        flush=lambda: None,
        commit=lambda: commits.append(True),
        rollback=lambda: rollbacks.append(True),
    )
    service.ecos.next_item = lambda: 10
    service.ecos.next_position = lambda: 20
    service.history.add = histories.append
    service.MAX_PREVIEW_ROWS = 1

    result = service.import_file(
        create_workbook(
            rows=[existing_row, new_row]
        ),
        "controle.xlsx",
    )

    assert result["imported_rows"] == 1
    assert result["duplicate_rows"] == 1
    assert result["imported"][0]["item"] == 10
    assert len(persisted) == 1
    assert persisted[0].product == "BBB"
    assert persisted[0].position == 20
    assert persisted[0].month == "JAN"
    assert histories
    assert commits == [True]
    assert rollbacks == []


def test_import_file_rolls_back_when_persistence_fails():

    service = create_service()
    persisted = []
    commits = []
    rollbacks = []

    def fail_flush():
        raise RuntimeError("flush failed")

    service.db = SimpleNamespace(
        add=persisted.append,
        flush=fail_flush,
        commit=lambda: commits.append(True),
        rollback=lambda: rollbacks.append(True),
    )
    service.ecos.next_item = lambda: 10
    service.ecos.next_position = lambda: 20

    with pytest.raises(RuntimeError, match="flush failed"):
        service.import_file(
            create_workbook(),
            "controle.xlsx",
        )

    assert len(persisted) == 1
    assert commits == []
    assert rollbacks == [True]


def test_preview_exact_existing_row_is_duplicate():

    baseline = create_service().preview(
        create_workbook(),
        "controle.xlsx",
    )
    parsed_data = baseline["rows"][0]["data"]
    existing_record = SimpleNamespace(
        **parsed_data,
        id="existing-id",
        item=55,
    )

    service = create_service(
        existing={
            "ECO-001": [existing_record],
        }
    )

    result = service.preview(
        create_workbook(),
        "controle.xlsx",
    )

    assert result["new_rows"] == 0
    assert result["update_rows"] == 0
    assert result["duplicate_rows"] == 1
    assert result["can_import"] is False
    assert result["rows"][0]["action"] == "DUPLICATE"
    assert result["rows"][0]["duplicate_source"] == "DATABASE"
    assert result["rows"][0]["existing"] == {
        "id": "existing-id",
        "item": 55,
    }


def test_preview_same_eco_with_different_data_is_new():

    baseline = create_service().preview(
        create_workbook(),
        "controle.xlsx",
    )
    existing_data = dict(
        baseline["rows"][0]["data"]
    )
    existing_data["product"] = "DIFFERENT-PRODUCT"
    existing_record = SimpleNamespace(
        **existing_data,
        id="existing-id",
        item=55,
    )

    service = create_service(
        existing={
            "ECO-001": [existing_record],
        }
    )

    result = service.preview(
        create_workbook(),
        "controle.xlsx",
    )

    assert result["new_rows"] == 1
    assert result["update_rows"] == 0
    assert result["duplicate_rows"] == 0
    assert result["rows"][0]["action"] == "NEW"
    assert result["rows"][0]["existing"] is None


def test_preview_mixed_new_and_duplicate_rows_can_import():

    new_row = {
        "item": 1,
        "product": "AAA",
        "obu": "NW1",
        "owner": "owner.test",
        "item_type": "MEC",
        "eco_type": "REGULAR",
        "change_bom": "YES",
        "status": "RELEASED",
        "receb": "NORMAL",
        "eco": "ECO-MIXED",
        "change_reason": "Alteração teste",
        "hq_eco_release_date": date(2026, 1, 2),
        "az_eco_register_date": date(2026, 1, 5),
    }
    baseline = create_service().preview(
        create_workbook(rows=[new_row]),
        "controle.xlsx",
    )
    existing_record = SimpleNamespace(
        **baseline["rows"][0]["data"],
        id="existing-id",
        item=55,
    )
    different_row = {
        **new_row,
        "item": 2,
        "product": "NEW-PRODUCT",
    }
    service = create_service(
        existing={
            "ECO-MIXED": [existing_record],
        }
    )

    result = service.preview(
        create_workbook(
            rows=[new_row, different_row]
        ),
        "controle.xlsx",
    )

    assert result["new_rows"] == 1
    assert result["duplicate_rows"] == 1
    assert result["error_rows"] == 0
    assert result["can_import"] is True
    assert result["rows"][0]["action"] == "DUPLICATE"
    assert result["rows"][1]["action"] == "NEW"


def test_preview_identical_rows_in_file_are_duplicate_except_item():

    original_row = {
        "item": 1,
        "product": "AAA",
        "obu": "NW1",
        "owner": "owner.test",
        "item_type": "MEC",
        "eco_type": "REGULAR",
        "change_bom": "YES",
        "status": "WORKING",
        "receb": "NORMAL",
        "eco": "ECO-FILE-DUPLICATE",
        "change_reason": "Alteração teste",
        "hq_eco_release_date": date(2026, 1, 2),
        "az_eco_register_date": date(2026, 1, 5),
    }
    repeated_row = {
        **original_row,
        "item": 2,
    }

    result = create_service().preview(
        create_workbook(
            rows=[original_row, repeated_row]
        ),
        "controle.xlsx",
    )

    assert result["new_rows"] == 1
    assert result["duplicate_rows"] == 1
    assert result["rows"][0]["action"] == "NEW"
    assert result["rows"][0]["duplicate_source"] is None
    assert result["rows"][1]["action"] == "DUPLICATE"
    assert result["rows"][1]["duplicate_source"] == "FILE"


# =========================================================
# CAMPOS CALCULADOS
# =========================================================

def test_preview_recalculates_fields():

    service = (
        create_service()
    )


    result = (
        service.preview(
            create_workbook(),
            "controle.xlsx",
        )
    )


    derived = (
        result[
            "rows"
        ][0][
            "derived"
        ]
    )


    assert (
        derived["gap"]
        ==
        "OK"
    )


    assert (
        derived[
            "delay_eco_register"
        ]
        == 3
    )


    assert (
        derived[
            "gap_start_az_eco"
        ]
        == 1
    )


    # 05/01/2026 até 12/01/2026
    # = 6 dias úteis.

    assert (
        derived["az_gap"]
        == 6
    )


    assert (
        derived[
            "contar_eco_7"
        ]
        is False
    )


# =========================================================
# ECO REPETIDA NO PRÓPRIO XLSX
# =========================================================

def test_preview_repeated_eco_is_allowed():

    rows = [
        {
            "item":
                1,

            "obu":
                "NW1",

            "owner":
                "owner.test",

            "item_type":
                "MEC",

            "eco_type":
                "REGULAR",

            "status":
                "WORKING",

            "eco":
                "ECO-REPETIDA",

            "az_eco_register_date":
                date(
                    2026,
                    1,
                    5,
                ),
        },

        {
            "item":
                2,

            "obu":
                "NWE",

            "owner":
                "owner.test",

            "item_type":
                "MEC",

            "eco_type":
                "REGULAR",

            "status":
                "WORKING",

            "eco":
                "ECO-REPETIDA",

            "az_eco_register_date":
                date(
                    2026,
                    1,
                    6,
                ),
        },
    ]


    service = create_service(
        obu_map={
            "NW1": "GLZ",
            "NWE": "GLZ",
        }
    )


    result = (
        service.preview(
            create_workbook(
                rows=rows
            ),
            "controle.xlsx",
        )
    )


    assert (
        result["total_rows"]
        == 2
    )


    assert (
        result["new_rows"]
        == 2
    )


    assert (
        result["error_rows"]
        == 0
    )


    assert (
        result["can_import"]
        is True
    )


    assert (
        result["rows"][0]["action"]
        ==
        "NEW"
    )


    assert (
        result["rows"][1]["action"]
        ==
        "NEW"
    )


# =========================================================
# OBU SEM MAPEAMENTO
# =========================================================

def test_preview_requires_obu_mapping():

    rows = [
        {
            "item":
                1,

            "obu":
                "NW9",

            "owner":
                "owner.test",

            "item_type":
                "MEC",

            "eco_type":
                "REGULAR",

            "status":
                "WORKING",

            "eco":
                "ECO-NW9",

            "az_eco_register_date":
                date(
                    2026,
                    1,
                    5,
                ),
        }
    ]


    service = (
        create_service(
            obu_map={
                "NW1":
                    "GLZ",
            }
        )
    )


    result = (
        service.preview(
            create_workbook(
                rows=rows
            ),
            "controle.xlsx",
        )
    )


    assert (
        result["error_rows"]
        == 1
    )


    assert (
        result["can_import"]
        is False
    )


    assert any(
        "OBU"
        in error

        for error
        in result[
            "rows"
        ][0][
            "errors"
        ]
    )


def test_preview_limits_to_errors_when_errors_exist():

    rows = [
        {
            "item": item,
            "product": "AAA",
            "obu": obu,
            "owner": "owner.test",
            "item_type": "MEC",
            "eco_type": "REGULAR",
            "status": "WORKING",
            "eco": f"ECO-{item}",
            "az_eco_register_date": date(2026, 1, 5),
        }
        for item, obu in (
            (1, "NW1"),
            (2, "NW9"),
            (3, "NW9"),
        )
    ]
    service = create_service()
    service.MAX_PREVIEW_ROWS = 1

    result = service.preview(
        create_workbook(rows=rows),
        "controle.xlsx",
    )

    assert result["total_rows"] == 3
    assert result["error_rows"] == 2
    assert result["preview_mode"] == "ERRORS"
    assert result["preview_rows_total"] == 2
    assert result["preview_rows_count"] == 1
    assert result["preview_truncated"] is True
    assert len(result["rows"]) == 1
    assert result["rows"][0]["action"] == "ERROR"
    assert result["rows"][0]["data"]["eco"] == "ECO-2"


# =========================================================
# OWNER SEM MAPEAMENTO
# =========================================================

def test_preview_requires_owner_mapping():

    rows = [
        {
            "item":
                1,

            "obu":
                "NW1",

            "owner":
                "owner.inexistente",

            "item_type":
                "MEC",

            "eco_type":
                "REGULAR",

            "status":
                "WORKING",

            "eco":
                "ECO-OWNER",

            "az_eco_register_date":
                date(
                    2026,
                    1,
                    5,
                ),
        }
    ]


    service = (
        create_service(
            owner_map={
                "owner.test":
                    "BOM",
            }
        )
    )


    result = (
        service.preview(
            create_workbook(
                rows=rows
            ),
            "controle.xlsx",
        )
    )


    assert (
        result["error_rows"]
        == 1
    )


    assert any(
        "Owner"
        in error

        for error
        in result[
            "rows"
        ][0][
            "errors"
        ]
    )


# =========================================================
# ABA OBRIGATÓRIA
# =========================================================

def test_preview_requires_ctrl_geral():

    service = (
        create_service()
    )


    with pytest.raises(
        HTTPException
    ) as exc:

        service.preview(
            create_workbook(
                sheet_name="OUTRA ABA"
            ),
            "controle.xlsx",
        )


    assert (
        exc.value.status_code
        == 422
    )


    assert (
        "CTRL GERAL"
        in str(
            exc.value.detail
        )
    )


# =========================================================
# EXTENSÃO
# =========================================================

def test_preview_rejects_non_xlsx():

    service = (
        create_service()
    )


    with pytest.raises(
        HTTPException
    ) as exc:

        service.preview(
            b"arquivo",
            "controle.csv",
        )


    assert (
        exc.value.status_code
        == 422
    )


# =========================================================
# ECO OBRIGATÓRIA
# =========================================================

def test_preview_requires_eco_code():

    rows = [
        {
            "item":
                1,

            "obu":
                "NW1",

            "owner":
                "owner.test",

            "item_type":
                "MEC",

            "eco_type":
                "REGULAR",

            "status":
                "WORKING",

            "eco":
                None,

            "az_eco_register_date":
                date(
                    2026,
                    1,
                    5,
                ),
        }
    ]


    service = (
        create_service()
    )


    result = (
        service.preview(
            create_workbook(
                rows=rows
            ),
            "controle.xlsx",
        )
    )


    assert (
        result["error_rows"]
        == 1
    )


    assert (
        result["rows"][0]["action"]
        ==
        "ERROR"
    )


    assert any(
        "ECO está vazia"
        in error

        for error
        in result[
            "rows"
        ][0][
            "errors"
        ]
    )

def edit_workbook(content, edit):
    workbook = load_workbook(io.BytesIO(content))
    try:
        edit(workbook.active)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


@pytest.mark.parametrize(("column", "field", "value"), [
    (37, "effective_date", date(2026, 2, 4)),
    (39, "agreement_start_1", date(2026, 2, 5)),
    (40, "agreement_finish_1", date(2026, 2, 6)),
    (42, "agreement_start_2", date(2026, 2, 7)),
    (43, "agreement_finish_2", date(2026, 2, 8)),
    (45, "agreement_start_3", date(2026, 2, 9)),
    (46, "agreement_finish_3", date(2026, 2, 10)),
    (49, "second_aprov_rd", "R&D-COLUMN-49"),
    (50, "second_aprov_rd_start_1", date(2026, 1, 8)),
    (51, "second_aprov_rd_finish_1", date(2026, 2, 12)),
    (53, "second_aprov_rd_start_2", date(2026, 2, 13)),
    (54, "second_aprov_rd_finish_2", date(2026, 2, 14)),
    (65, "change_reason_az", "Reason column 65"),
    (66, "model_az", "Model column 66"),
    (67, "new_model", "YES"),
    (68, "origem_approval", "HQ"),
    (69, "event", "YES"),
    (70, "comments", "Comments column 70"),
    (71, "set_eco", "SET-COLUMN-71"),
    (72, "set_eco_register_date", date(2026, 2, 15)),
    (73, "set_eco_release_date", date(2026, 2, 16)),
    (74, "set_eco_change_reason", "Reason column 74"),
    (75, "set_eco_model", "Model column 75"),
])
def test_preview_reads_official_column_positions(column, field, value):
    content = edit_workbook(create_workbook(), lambda sheet: sheet.cell(8, column, value))
    result = create_service().preview(content, "controle.xlsx")
    assert result["error_rows"] == 0
    assert result["rows"][0]["data"][field] == value


@pytest.mark.parametrize("header", ["Effective Date", "effective date", "  Effective   Date  ", "Effective\nDate"])
def test_effective_date_header_normalization(header):
    content = edit_workbook(create_workbook(), lambda sheet: sheet.cell(7, 37, header))
    result = create_service().preview(content, "controle.xlsx")
    assert result["rows"][0]["data"]["effective_date"] == date(2026, 2, 4)


def test_effective_date_rejects_wrong_header():
    content = edit_workbook(create_workbook(), lambda sheet: sheet.cell(7, 37, "Wrong date"))
    with pytest.raises(HTTPException) as exc:
        create_service().preview(content, "controle.xlsx")
    assert exc.value.status_code == 422


def test_old_ctrl_geral_layout_is_rejected():
    def old_layout(sheet):
        sheet.delete_cols(61)
        sheet.delete_cols(37)
    content = edit_workbook(create_workbook(), old_layout)
    with pytest.raises(HTTPException) as exc:
        create_service().preview(content, "controle.xlsx")
    assert exc.value.status_code == 422


def test_effective_date_formula_is_rejected():
    content = edit_workbook(create_workbook(), lambda sheet: sheet.cell(8, 37, "=TODAY()"))
    result = create_service().preview(content, "controle.xlsx")
    assert result["can_import"] is False
    assert any("EFFECTIVE DATE" in error for error in result["rows"][0]["errors"])


@pytest.mark.parametrize("calculated_value", [999, "=999"])
def test_gap_14d_is_not_imported_or_persisted(calculated_value):
    specs = EcoImportService.COLUMN_SPECS
    assert 61 not in specs
    assert "gap_14d" not in {spec["field"] for spec in specs.values()}
    calculated_columns = set(range(1, 76)) - set(specs) - {2}
    def populate_calculated(sheet):
        for column in calculated_columns:
            sheet.cell(8, column, calculated_value)
    content = edit_workbook(create_workbook(), populate_calculated)
    service = create_service()
    preview = service.preview(content, "controle.xlsx")
    assert preview["error_rows"] == 0
    assert "gap_14d" not in preview["rows"][0]["data"]
    assert calculated_value not in preview["rows"][0]["data"].values()
    persisted = []
    service.db = SimpleNamespace(
        add=persisted.append, flush=lambda: None,
        commit=lambda: None, rollback=lambda: None,
    )
    service.ecos.next_item = lambda: 10
    service.ecos.next_position = lambda: 20
    service.history.add = lambda event: None
    result = service.import_file(content, "controle.xlsx")
    assert result["imported_rows"] == 1
    assert persisted[0].effective_date == date(2026, 2, 4)
    assert persisted[0].gap_14d_note is None
    assert not hasattr(persisted[0], "gap_14d")
    assert all(getattr(persisted[0], column.name) != calculated_value
               for column in persisted[0].__table__.columns)
