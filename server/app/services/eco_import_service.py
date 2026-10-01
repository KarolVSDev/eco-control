import io
import re
import unicodedata

from datetime import (
    date,
    datetime,
)

from types import (
    SimpleNamespace,
)

from fastapi import (
    HTTPException,
)

from openpyxl import (
    load_workbook,
)

from openpyxl.utils.datetime import (
    from_excel,
)

from sqlalchemy import (
    func,
    select,
)

from app.models.entities import (
    AnalystPermission,
    Eco,
    EcoHistory,
)

from app.repository.eco_repository import (
    EcoRepository,
)

from app.repository.history_repository import (
    HistoryRepository,
)

from app.repository.settings_repository import (
    SettingsRepository,
)

from app.utils.eco_calculations import (
    compute_eco_fields,
)


class EcoImportService:

    # =========================================================
    # CONFIGURAÇÃO
    # =========================================================

    SHEET_NAME = "CTRL GERAL"

    INTERNAL_MONTHS = [
        "JAN",
        "FEB",
        "MAR",
        "APR",
        "MAY",
        "JUN",
        "JUL",
        "AUG",
        "SEP",
        "OCT",
        "NOV",
        "DEC",
    ]

    MAX_SCAN_ROWS = 10000

    MAX_PREVIEW_ROWS = 1000

    EMPTY_ROWS_TO_STOP = 100


    # =========================================================
    # COLUNAS DA PLANILHA
    #
    # Somente campos de origem são importados.
    #
    # Fórmulas como:
    #
    # GAP
    # MONTH
    # AU
    # GROUP
    # ECO Origem
    # AZ Gap
    # Release Month
    # etc.
    #
    # NÃO são importadas.
    # =========================================================

    COLUMN_SPECS = {
        # -----------------------------------------------------
        # IDENTIFICAÇÃO
        # -----------------------------------------------------

        4: {
            "field": "product",
            "header": "PRODUCT",
            "type": "text",
        },

        6: {
            "field": "obu",
            "header": "OBU",
            "type": "text",
        },

        # -----------------------------------------------------
        # CLASSIFICAÇÃO
        # -----------------------------------------------------

        8: {
            "field": "owner",
            "header": "OWNER",
            "type": "text",
        },

        9: {
            "field": "item_type",
            "header": "ITEM TYPE",
            "type": "text",
        },

        10: {
            "field": "eco_type",
            "header": "ECO TYPE",
            "type": "text",
        },

        11: {
            "field": "change_bom",
            "header": "CHANGE BOM?",
            "type": "text",
        },

        12: {
            "field": "status",
            "header": "STATUS",
            "type": "text",
        },

        13: {
            "field": "receb",
            "header": "RECEB.",
            "type": "text",
        },

        # -----------------------------------------------------
        # ECO HQ
        # -----------------------------------------------------

        14: {
            "field": "eco",
            "header": "ECO",
            "type": "text",
        },

        15: {
            "field": "change_reason",
            "header": "CHANGE REASON",
            "type": "text",
        },

        16: {
            "field": "hq_eco_release_date",
            "header": "HQ ECO RELEASE DATE",
            "type": "date",
        },

        17: {
            "field": "az_eco_register_date",
            "header": "AZ ECO REGISTER DATE",
            "type": "date",
        },

        18: {
            "field": "auto_ecr",
            "header": "AUTO ECR?",
            "type": "text",
        },

        21: {
            "field": "council_meeting_week",
            "header": "COUNCIL MEETING WEEK",
            "type": "text",
        },

        # -----------------------------------------------------
        # ORIGEM 1
        # -----------------------------------------------------

        22: {
            "field": "hz_in_kr_eco_1",
            "header": "HZ/IN/KR ECO(1)",
            "type": "text",
        },

        23: {
            "field": "hz_sh_in_receive_date_1",
            "header": "HZ/IN/KR RECEIVE DATE (1)",
            "type": "date",
        },

        # -----------------------------------------------------
        # ORIGEM 2
        # -----------------------------------------------------

        25: {
            "field": "hz_in_kr_eco_2",
            "header": "HZ/IN/KR ECO(2)",
            "type": "text",
        },

        26: {
            "field": "hz_sh_in_receive_date_2",
            "header": "HZ/SH/IN RECEIVE DATE (2)",
            "type": "date",
        },

        # -----------------------------------------------------
        # ORIGEM 3
        # -----------------------------------------------------

        28: {
            "field": "hz_in_kr_eco_3",
            "header": "HZ/IN/KR ECO(3)",
            "type": "text",
        },

        29: {
            "field": "hz_sh_in_receive_date_3",
            "header": "HZ/SH/IN RECEIVE DATE (3)",
            "type": "date",
        },

        # -----------------------------------------------------
        # AZ ECO
        # -----------------------------------------------------

        31: {
            "field": "az_eco_no",
            "header": "AZ_ECO_NO.",
            "type": "text",
        },

        32: {
            "field": "change_reason2",
            "header": "CHANGE REASON2",
            "type": "text",
        },

        33: {
            "field": "change_bom_az",
            "header": "CHANGE BOM?",
            "type": "text",
        },

        34: {
            "field": "az_eco_creation_date",
            "header": "AZ ECO CREATION DATE",
            "type": "date",
        },

        # -----------------------------------------------------
        # AGREEMENT OTHER DEPTS
        # -----------------------------------------------------

        38: {
            "field": "agreement_start_1",
            "header":
                "AGREEMENT OTHER DEPTS START (1)",
            "type": "date",
        },

        39: {
            "field": "agreement_finish_1",
            "header":
                "AGREEMENT OTHER DEPTS FINSHI (1)",
            "type": "date",
        },

        41: {
            "field": "agreement_start_2",
            "header":
                "AGREEMENT OTHER DEPTS START (2)",
            "type": "date",
        },

        42: {
            "field": "agreement_finish_2",
            "header":
                "AGREEMENT OTHER DEPTS FINSHI (2)",
            "type": "date",
        },

        44: {
            "field": "agreement_start_3",
            "header":
                "AGREEMENT OTHER DEPTS START (3)",
            "type": "date",
        },

        45: {
            "field": "agreement_finish_3",
            "header":
                "AGREEMENT OTHER DEPTS FINSHI (3)",
            "type": "date",
        },

                # -----------------------------------------------------
        # R&D APPROVAL
        # -----------------------------------------------------

        48: {
            "field": "second_aprov_rd",
            "header": "2ST APROV R&D",
            "type": "text",
        },

        49: {
            "field": "second_aprov_rd_start_1",
            "header":
                "2ST APROV R&D START (1)",
            "type": "date",
        },

        50: {
            "field": "second_aprov_rd_finish_1",
            "header":
                "2ST APROV R&D FINISH (1)",
            "type": "date",
        },

        52: {
            "field": "second_aprov_rd_start_2",
            "header":
                "2ST APROV R&D START (2)",
            "type": "date",
        },

        53: {
            "field": "second_aprov_rd_finish_2",
            "header":
                "2ST APROV R&D FINISH (2)",
            "type": "date",
        },
        # -----------------------------------------------------
        # ADICIONAIS
        # -----------------------------------------------------

        63: {
            "field": "change_reason_az",
            "header": "CHANGE REASON",
            "type": "text",
        },

        64: {
            "field": "model_az",
            "header": "MODEL",
            "type": "text",
        },

        65: {
            "field": "new_model",
            "header":
                "NEW MODEL? (CHECK IN NPI PLAN 26Y)",
            "type": "text",
        },

        66: {
            "field": "origem_approval",
            "header": "ORIGEM APPROVAL",
            "type": "text",
        },

        67: {
            "field": "event",
            "header": "EVENT",
            "type": "text",
        },

        68: {
            "field": "comments",
            "header": "COMMENTS",
            "type": "text",
        },

        # -----------------------------------------------------
        # SET ECO
        # -----------------------------------------------------

        69: {
            "field": "set_eco",
            "header": "SET ECO",
            "type": "text",
        },

        70: {
            "field": "set_eco_register_date",
            "header": "SET ECO REGISTER DATE",
            "type": "date",
        },

        71: {
            "field": "set_eco_release_date",
            "header": "SET ECO RELEASE DATE",
            "type": "date",
        },

        72: {
            "field": "set_eco_change_reason",
            "header": "CHANGE REASON",
            "type": "text",
        },

        73: {
            "field": "set_eco_model",
            "header": "MODEL",
            "type": "text",
        },
    }


    # =========================================================
    # CAMPOS YES / NO
    # =========================================================

    YES_NO_FIELDS = {
        "change_bom",
        "auto_ecr",
        "change_bom_az",
        "new_model",
        "event",
    }


    # =========================================================
    # OPÇÕES
    # =========================================================

    ITEM_TYPE_OPTIONS = {
        "MEC",
        "DEV",
        "ELET",
    }


    ECO_TYPE_OPTIONS = {
        "LOCAL",
        "TOOL",
        "TEMP",
        "DC-ECO",
        "SOFTWARE",
        "REGULAR",
        "DESENHO - IMPRESSOS/BOX",
        "DESENHO - NÃO APLICADA A AZ",
        "DESENHO - ENVIADA E-MAIL",
        "DESENHO",
    }


    STATUS_OPTIONS = {
        "WORKING",
        "ON HOLD",
        "WAITING HZ/IN/ND",
        "MEC/HW ANALISYS",
        "COUNCIL MEETING",
        "PROCESSING",
        "SPOC ON APPROVAL",
        "WAITING NEW ECO TO FIX BOM",
        "REJECTED",
        "RELEASED",
        "TO BE CANCELLED",
        "CANCELLED",
    }


    RECEB_OPTIONS = {
        "NORMAL",
        "MISSING 1",
        "MISSING 2",
    }


    ORIGEM_APPROVAL_OPTIONS = {
        "HQ",
        "ND",
        "RC",
        "HZ",
        "HQ/HZ",
        "IN",
        "MISSING",
        "NA",
    }


    # =========================================================
    # CONSTRUTOR
    # =========================================================

    def __init__(
        self,
        db,
        user,
    ):
        self.db = db
        self.user = user

        self.settings = (
            SettingsRepository(
                db
            )
        )

        self.ecos = (
            EcoRepository(
                db
            )
        )

        self.history = (
            HistoryRepository(
                db
            )
        )


    # =========================================================
    # PERMISSÃO
    # =========================================================

    def _can_bulk_edit(
        self,
    ):
        if (
            self.user.role
            ==
            "admin"
        ):
            return True


        permission = (
            self.db.scalar(
                select(
                    AnalystPermission
                ).where(
                    AnalystPermission.user_email
                    ==
                    self.user.email
                )
            )
        )


        return bool(
            permission
            and
            permission.can_bulk_edit
        )


    def _ensure_permission(
        self,
    ):
        if self._can_bulk_edit():
            return


        raise HTTPException(
            status_code=403,
            detail=(
                "Você não possui permissão "
                "para importar ECOs em massa."
            ),
        )


    # =========================================================
    # NORMALIZAÇÃO
    # =========================================================

    @staticmethod
    def _normalize_header(
        value,
    ):
        if value is None:
            return ""


        text = str(
            value
        )


        text = text.replace(
            "\n",
            " ",
        )


        text = text.replace(
            "\r",
            " ",
        )


        return (
            " ".join(
                text.split()
            )
            .strip()
            .upper()
        )


    @staticmethod
    def _option_key(
        value,
    ):
        if value is None:
            return ""


        text = (
            str(value)
            .strip()
            .upper()
        )


        text = (
            unicodedata
            .normalize(
                "NFKD",
                text,
            )
        )


        text = "".join(
            char
            for char
            in text
            if not unicodedata.combining(
                char
            )
        )


        text = re.sub(
            r"\s*/\s*",
            "/",
            text,
        )


        text = re.sub(
            r"\s*-\s*",
            " - ",
            text,
        )


        text = " ".join(
            text.split()
        )


        return text


    @staticmethod
    def _clean_text(
        value,
    ):
        if value is None:
            return None


        if isinstance(
            value,
            bool,
        ):
            return (
                "YES"
                if value
                else "NO"
            )


        text = str(
            value
        ).strip()


        if (
            not text
            or
            text
            in {
                "-",
                "—",
            }
        ):
            return None


        return text


    # =========================================================
    # CANONICALIZAÇÃO DAS OPÇÕES
    # =========================================================

    def _canonical_option(
        self,
        value,
        options,
    ):
        cleaned = (
            self._clean_text(
                value
            )
        )


        if cleaned is None:
            return None


        expected_key = (
            self._option_key(
                cleaned
            )
        )


        for option in options:

            if (
                self._option_key(
                    option
                )
                ==
                expected_key
            ):
                return option


        return cleaned


    def _normalize_field_value(
        self,
        field,
        value,
    ):
        cleaned = (
            self._clean_text(
                value
            )
        )


        if cleaned is None:
            return None


        # -----------------------------------------------------
        # YES / NO
        # -----------------------------------------------------

        if (
            field
            in self.YES_NO_FIELDS
        ):
            key = (
                self._option_key(
                    cleaned
                )
            )


            if key in {
                "YES",
                "SIM",
                "TRUE",
                "1",
            }:
                return "YES"


            if key in {
                "NO",
                "NAO",
                "FALSE",
                "0",
            }:
                return "NO"


            return cleaned


        # -----------------------------------------------------
        # STATUS
        # -----------------------------------------------------

        if field == "status":

            return (
                self._canonical_option(
                    cleaned,
                    self.STATUS_OPTIONS,
                )
            )


        # -----------------------------------------------------
        # ITEM TYPE
        # -----------------------------------------------------

        if field == "item_type":

            return (
                self._canonical_option(
                    cleaned,
                    self.ITEM_TYPE_OPTIONS,
                )
            )


        # -----------------------------------------------------
        # ECO TYPE
        # -----------------------------------------------------

        if field == "eco_type":

            return (
                self._canonical_option(
                    cleaned,
                    self.ECO_TYPE_OPTIONS,
                )
            )


        # -----------------------------------------------------
        # RECEB
        # -----------------------------------------------------

        if field == "receb":

            return (
                self._canonical_option(
                    cleaned,
                    self.RECEB_OPTIONS,
                )
            )


        # -----------------------------------------------------
        # ORIGEM APPROVAL
        # -----------------------------------------------------

        if (
            field
            ==
            "origem_approval"
        ):

            return (
                self._canonical_option(
                    cleaned,
                    self.ORIGEM_APPROVAL_OPTIONS,
                )
            )


        # -----------------------------------------------------
        # OBU
        # -----------------------------------------------------

        if field == "obu":

            return cleaned.upper()


        # -----------------------------------------------------
        # ECO
        # -----------------------------------------------------

        if field == "eco":

            return cleaned.upper()


        return cleaned


    # =========================================================
    # VALIDAÇÃO DE OPÇÃO
    # =========================================================

    def _validate_option(
        self,
        field,
        value,
        errors,
    ):
        if value is None:
            return


        options = None


        if field in self.YES_NO_FIELDS:

            options = {
                "YES",
                "NO",
            }


        elif field == "item_type":

            options = (
                self.ITEM_TYPE_OPTIONS
            )


        elif field == "eco_type":

            options = (
                self.ECO_TYPE_OPTIONS
            )


        elif field == "status":

            options = (
                self.STATUS_OPTIONS
            )


        elif field == "receb":

            options = (
                self.RECEB_OPTIONS
            )


        elif (
            field
            ==
            "origem_approval"
        ):

            options = (
                self.ORIGEM_APPROVAL_OPTIONS
            )


        if options is None:
            return


        if value not in options:

            errors.append(
                (
                    f"Valor inválido para "
                    f"{field}: {value}"
                )
            )


    # =========================================================
    # DATAS
    # =========================================================

    def _parse_date(
        self,
        value,
        workbook_epoch,
    ):
        if value is None:
            return None


        if isinstance(
            value,
            datetime,
        ):
            return value.date()


        if isinstance(
            value,
            date,
        ):
            return value


        if isinstance(
            value,
            (
                int,
                float,
            ),
        ):
            try:

                parsed = from_excel(
                    value,
                    epoch=(
                        workbook_epoch
                    ),
                )


                if isinstance(
                    parsed,
                    datetime,
                ):
                    return parsed.date()


                if isinstance(
                    parsed,
                    date,
                ):
                    return parsed


            except Exception:

                return None


        if isinstance(
            value,
            str,
        ):

            text = (
                value
                .strip()
            )


            if not text:
                return None


            formats = [
                "%Y-%m-%d",
                "%d/%m/%Y",
                "%d/%m/%y",
                "%m/%d/%Y",
            ]


            for date_format in formats:

                try:

                    return (
                        datetime.strptime(
                            text,
                            date_format,
                        )
                        .date()
                    )


                except ValueError:

                    continue


        return None


    # =========================================================
    # ENCONTRAR CABEÇALHO
    # =========================================================

    def _find_header_row(
        self,
        worksheet,
    ):
        """
        Aceita:

        - planilha original, onde o header está
          normalmente na linha 7;

        - XLSX exportado pelo próprio sistema,
          onde o header fica na linha 2.

        Não dependemos de um número fixo.
        """

        max_row = 30


        for row_number in range(
            1,
            max_row + 1,
        ):

            gap = (
                self._normalize_header(
                    worksheet.cell(
                        row=row_number,
                        column=1,
                    ).value
                )
            )


            item = (
                self._normalize_header(
                    worksheet.cell(
                        row=row_number,
                        column=2,
                    ).value
                )
            )


            eco = (
                self._normalize_header(
                    worksheet.cell(
                        row=row_number,
                        column=14,
                    ).value
                )
            )


            register_date = (
                self._normalize_header(
                    worksheet.cell(
                        row=row_number,
                        column=17,
                    ).value
                )
            )


            if (
                gap == "GAP"
                and
                item == "ITEM"
                and
                eco == "ECO"
                and
                register_date
                ==
                "AZ ECO REGISTER DATE"
            ):

                return row_number


        raise HTTPException(
            status_code=422,
            detail=(
                "Não foi possível localizar "
                "o cabeçalho da aba CTRL GERAL."
            ),
        )


    # =========================================================
    # VALIDAR ESTRUTURA
    # =========================================================

    def _validate_headers(
        self,
        worksheet,
        header_row,
    ):
        invalid = []


        for (
            column,
            spec,
        ) in self.COLUMN_SPECS.items():

            actual = (
                self._normalize_header(
                    worksheet.cell(
                        row=header_row,
                        column=column,
                    ).value
                )
            )


            expected = (
                self._normalize_header(
                    spec["header"]
                )
            )


            # Alguns títulos da planilha possuem
            # observações complementares.

            if (
                expected
                not in actual
                and
                actual
                not in expected
            ):

                invalid.append(
                    {
                        "column":
                            column,

                        "field":
                            spec[
                                "field"
                            ],

                        "expected":
                            spec[
                                "header"
                            ],

                        "found":
                            (
                                actual
                                or
                                "(vazio)"
                            ),
                    }
                )


        if invalid:

            details = "; ".join(
                (
                    f"coluna {item['column']} "
                    f"esperada '{item['expected']}', "
                    f"encontrada '{item['found']}'"
                )
                for item
                in invalid[:8]
            )


            raise HTTPException(
                status_code=422,
                detail=(
                    "A estrutura da aba "
                    "CTRL GERAL não corresponde "
                    "ao modelo esperado. "
                    f"{details}"
                ),
            )


    # =========================================================
    # LINHA VAZIA
    # =========================================================

    def _row_has_source_data(
        self,
        cells,
    ):
        # ITEM
        if (
            len(cells) >= 2
            and
            cells[1].value
            not in {
                None,
                "",
            }
        ):
            return True


        for column in self.COLUMN_SPECS:

            index = (
                column - 1
            )


            if (
                index
                >= len(cells)
            ):
                continue


            value = (
                cells[
                    index
                ].value
            )


            if value not in {
                None,
                "",
            }:

                return True


        return False


    # =========================================================
    # ITEM DA PLANILHA
    # =========================================================

    @staticmethod
    def _source_item(
        value,
    ):
        if value is None:
            return None


        if isinstance(
            value,
            bool,
        ):
            return None


        if isinstance(
            value,
            int,
        ):
            return value


        if isinstance(
            value,
            float,
        ):
            if value.is_integer():
                return int(
                    value
                )

            return None


        try:

            return int(
                str(value)
                .strip()
            )


        except (
            TypeError,
            ValueError,
        ):

            return None


    # =========================================================
    # PARSE DA LINHA
    # =========================================================

    def _parse_row(
        self,
        cells,
        row_number,
        workbook_epoch,
    ):
        data = {}

        errors = []

        warnings = []


        source_item = None


        if len(cells) >= 2:

            source_item = (
                self._source_item(
                    cells[1].value
                )
            )


        # -----------------------------------------------------
        # CAMPOS DE ORIGEM
        # -----------------------------------------------------

        for (
            column,
            spec,
        ) in self.COLUMN_SPECS.items():

            index = (
                column - 1
            )


            if (
                index
                >= len(cells)
            ):
                value = None

            else:
                cell = cells[
                    index
                ]

                value = (
                    cell.value
                )


                # Campo de origem não deve conter fórmula.

                if (
                    cell.data_type
                    ==
                    "f"
                ):

                    errors.append(
                        (
                            f"{spec['header']} "
                            "não pode conter fórmula."
                        )
                    )

                    value = None


            # -------------------------------------------------
            # DATA
            # -------------------------------------------------

            if (
                spec["type"]
                ==
                "date"
            ):

                if value in {
                    None,
                    "",
                    "-",
                }:

                    parsed = None

                else:

                    parsed = (
                        self._parse_date(
                            value,
                            workbook_epoch,
                        )
                    )


                    if (
                        parsed
                        is None
                    ):

                        errors.append(
                            (
                                f"Data inválida em "
                                f"{spec['header']}: "
                                f"{value}"
                            )
                        )


                data[
                    spec["field"]
                ] = parsed


            # -------------------------------------------------
            # TEXTO
            # -------------------------------------------------

            else:

                normalized = (
                    self._normalize_field_value(
                        spec[
                            "field"
                        ],
                        value,
                    )
                )


                self._validate_option(
                    spec[
                        "field"
                    ],
                    normalized,
                    errors,
                )


                data[
                    spec[
                        "field"
                    ]
                ] = normalized


        # -----------------------------------------------------
        # ECO OBRIGATÓRIA
        # -----------------------------------------------------

        eco_code = (
            data.get(
                "eco"
            )
        )


        if not eco_code:

            errors.append(
                "A coluna ECO está vazia."
            )


        # -----------------------------------------------------
        # CONSISTÊNCIA DE DATAS
        # -----------------------------------------------------

        self._validate_date_pairs(
            data,
            errors,
            warnings,
        )


        return {
            "excel_row":
                row_number,

            "source_item":
                source_item,

            "eco":
                eco_code,

            "data":
                data,

            "errors":
                errors,

            "warnings":
                warnings,
        }


    # =========================================================
    # VALIDAR DATAS
    # =========================================================

    def _validate_date_pairs(
        self,
        data,
        errors,
        warnings,
    ):
        pairs = [
            (
                "agreement_start_1",
                "agreement_finish_1",
                "Agreement 1",
            ),

            (
                "agreement_start_2",
                "agreement_finish_2",
                "Agreement 2",
            ),

            (
                "agreement_start_3",
                "agreement_finish_3",
                "Agreement 3",
            ),

            (
                "second_aprov_rd_start_1",
                "second_aprov_rd_finish_1",
                "R&D 1",
            ),

            (
                "second_aprov_rd_start_2",
                "second_aprov_rd_finish_2",
                "R&D 2",
            ),

            (
                "set_eco_register_date",
                "set_eco_release_date",
                "SET ECO",
            ),
        ]


        for (
            start_key,
            finish_key,
            label,
        ) in pairs:

            start = data.get(
                start_key
            )

            finish = data.get(
                finish_key
            )


            if (
                start
                and
                not finish
            ):

                warnings.append(
                    (
                        f"{label} possui data "
                        "de início, mas não possui "
                        "data de conclusão."
                    )
                )


            if (
                finish
                and
                not start
            ):

                warnings.append(
                    (
                        f"{label} possui data "
                        "de conclusão, mas não "
                        "possui data de início."
                    )
                )


            if (
                start
                and
                finish
                and
                finish < start
            ):

                errors.append(
                    (
                        f"{label}: data final "
                        "anterior à data inicial."
                    )
                )


        hq_release = (
            data.get(
                "hq_eco_release_date"
            )
        )


        az_register = (
            data.get(
                "az_eco_register_date"
            )
        )


        if (
            hq_release
            and
            az_register
            and
            az_register < hq_release
        ):

            warnings.append(
                (
                    "AZ ECO REGISTER DATE é "
                    "anterior a HQ ECO RELEASE DATE."
                )
            )


        az_creation = (
            data.get(
                "az_eco_creation_date"
            )
        )


        if (
            az_register
            and
            az_creation
            and
            az_creation < az_register
        ):

            warnings.append(
                (
                    "AZ ECO CREATION DATE é "
                    "anterior a AZ ECO REGISTER DATE."
                )
            )


    # =========================================================
    # SETTINGS
    # =========================================================

    def _settings_maps(
        self,
    ):
        obu_map = {
            str(
                item.obu
            )
            .strip()
            .upper():
                item.au

            for item
            in self.settings
            .all_obu_au()
        }


        owner_map = {
            str(
                item.owner
            )
            .strip()
            .casefold():
                item.group_name

            for item
            in self.settings
            .all_owner_group()
        }


        return (
            obu_map,
            owner_map,
        )


    # =========================================================
    # REGISTROS JÁ EXISTENTES
    # =========================================================

    def _existing_rows_by_eco(
        self,
        eco_codes,
    ):
        """
        Busca todas as linhas existentes com ECOs
        presentes no XLSX, agrupadas pelo código ECO.
        """
        normalized_codes = {
            str(code).strip().upper()
            for code in eco_codes
            if code
        }

        if not normalized_codes:
            return {}

        rows = (
            self.db.scalars(
                select(Eco).where(
                    func.upper(Eco.eco).in_(normalized_codes)
                )
            )
            .all()
        )

        grouped = {}
        for eco in rows:
            if not eco.eco:
                continue

            key = str(eco.eco).strip().upper()
            grouped.setdefault(key, []).append(eco)

        return grouped


    def _normalize_existing_value(
        self,
        field,
        value,
    ):
        """Normalize database values as spreadsheet values are normalized."""
        spec = next(
            (
                candidate
                for candidate in self.COLUMN_SPECS.values()
                if candidate["field"] == field
            ),
            None,
        )

        if spec is None:
            return value

        if spec["type"] == "date":
            if isinstance(value, datetime):
                return value.date()
            return value

        return self._normalize_field_value(field, value)


    def _find_exact_existing(
        self,
        parsed_row,
        existing_rows,
    ):
        """Return an existing record only when every source field matches."""
        data = parsed_row["data"]

        for existing in existing_rows:
            is_same = True

            for spec in self.COLUMN_SPECS.values():
                field = spec["field"]
                spreadsheet_value = data.get(field)
                database_value = self._normalize_existing_value(
                    field,
                    getattr(existing, field, None),
                )

                if spreadsheet_value != database_value:
                    is_same = False
                    break

            if is_same:
                return existing

        return None

    # =========================================================
    # ADICIONAR CAMPOS DERIVADOS
    # =========================================================

    def _apply_derived_values(
        self,
        parsed_row,
        obu_map,
        owner_map,
    ):
        data = parsed_row[
            "data"
        ]


        errors = parsed_row[
            "errors"
        ]


        warnings = parsed_row[
            "warnings"
        ]


        obu = data.get(
            "obu"
        )


        owner = data.get(
            "owner"
        )


        # -----------------------------------------------------
        # AU
        # -----------------------------------------------------

        au = None


        if obu:

            au = (
                obu_map.get(
                    obu.upper()
                )
            )


            if au is None:

                errors.append(
                    (
                        f"OBU '{obu}' não possui "
                        "mapeamento OBU → AU "
                        "em Settings."
                    )
                )


        # -----------------------------------------------------
        # GROUP
        # -----------------------------------------------------

        group = None


        if owner:

            group = (
                owner_map.get(
                    owner.casefold()
                )
            )


            if group is None:

                errors.append(
                    (
                        f"Owner '{owner}' não possui "
                        "mapeamento Owner → Group "
                        "em Settings."
                    )
                )


        # -----------------------------------------------------
        # MONTH
        # -----------------------------------------------------

        register_date = (
            data.get(
                "az_eco_register_date"
            )
        )


        month = (
            self.INTERNAL_MONTHS[
                register_date.month
                - 1
            ]

            if register_date

            else None
        )


        if (
            register_date is None
        ):

            warnings.append(
                (
                    "MONTH não pôde ser calculado "
                    "porque AZ ECO REGISTER DATE "
                    "está vazia."
                )
            )


        # -----------------------------------------------------
        # CAMPOS CALCULADOS
        # -----------------------------------------------------

        calculation_row = (
            SimpleNamespace(
                **data
            )
        )


        computed = (
            compute_eco_fields(
                calculation_row
            )
        )


        parsed_row[
            "derived"
        ] = {
            "month":
                month,

            "au":
                au,

            "group":
                group,

            **computed,
        }


    # =========================================================
    # PREVIEW
    # =========================================================

    def _row_signature(
        self,
        parsed_row,
    ):
        """Build a signature from importable source fields only."""
        data = parsed_row["data"]

        return tuple(
            (
                spec["field"],
                data.get(spec["field"]),
            )
            for spec in self.COLUMN_SPECS.values()
        )


    # =========================================================
    # HISTÓRICO
    # =========================================================

    @staticmethod
    def _history_value(
        value,
    ):
        if value is None:
            return None


        if hasattr(
            value,
            "isoformat",
        ):

            return value.isoformat()


        return str(
            value
        )


    @staticmethod
    def _history_label(
        field,
    ):
        if field == "group_name":
            return "GROUP"


        return (
            field
            .replace(
                "_",
                " ",
            )
            .upper()
        )


    def _record_import_history(
        self,
        eco,
        data,
    ):
        """
        Registra no histórico todos os campos
        efetivamente preenchidos durante a importação.
        """

        fields = {
            "item":
                eco.item,

            "position":
                eco.position,

            "month":
                eco.month,

            "au":
                eco.au,

            "group_name":
                eco.group_name,
        }


        for (
            key,
            value,
        ) in data.items():

            if value is not None:

                fields[
                    key
                ] = value


        for (
            field,
            value,
        ) in fields.items():

            if value is None:
                continue


            self.history.add(
                EcoHistory(
                    eco_id=eco.id,

                    eco_code=eco.eco,

                    item=eco.item,

                    field_key=field,

                    field_label=(
                        self._history_label(
                            field
                        )
                    ),

                    old_value=None,

                    new_value=(
                        self._history_value(
                            value
                        )
                    ),

                    user_email=(
                        self.user.email
                    ),

                    user_name=(
                        self.user.full_name
                    ),

                    action="created",
                )
            )


    def preview(
        self,
        file_bytes,
        filename,
        _include_all_rows=False,
    ):
        """
        Analisa a planilha sem alterar o banco.

        Retorna:

        - linhas novas;
        - ECOs já existentes;
        - erros;
        - avisos;
        - valores derivados/calculados.

        Nenhum INSERT ou UPDATE ocorre aqui.
        """

        self._ensure_permission()


        # -----------------------------------------------------
        # ARQUIVO
        # -----------------------------------------------------

        if not filename:

            raise HTTPException(
                status_code=422,
                detail=(
                    "Nome do arquivo não informado."
                ),
            )


        if not (
            filename
            .lower()
            .endswith(
                ".xlsx"
            )
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    "Formato inválido. "
                    "Envie um arquivo .xlsx."
                ),
            )


        if not file_bytes:

            raise HTTPException(
                status_code=422,
                detail=(
                    "O arquivo enviado está vazio."
                ),
            )


        # -----------------------------------------------------
        # ABRIR WORKBOOK
        # -----------------------------------------------------

        try:

            workbook = (
                load_workbook(
                    filename=io.BytesIO(
                        file_bytes
                    ),
                    read_only=True,
                    data_only=False,
                )
            )


        except Exception as exc:

            raise HTTPException(
                status_code=422,
                detail=(
                    "Não foi possível abrir "
                    "o arquivo XLSX. "
                    f"Detalhe: {exc}"
                ),
            )


        try:

            # -------------------------------------------------
            # ABA
            # -------------------------------------------------

            if (
                self.SHEET_NAME
                not in
                workbook.sheetnames
            ):

                raise HTTPException(
                    status_code=422,
                    detail=(
                        "A planilha precisa possuir "
                        "a aba 'CTRL GERAL'."
                    ),
                )


            worksheet = workbook[
                self.SHEET_NAME
            ]


            # -------------------------------------------------
            # CABEÇALHO
            # -------------------------------------------------

            header_row = (
                self._find_header_row(
                    worksheet
                )
            )


            self._validate_headers(
                worksheet,
                header_row,
            )


            # -------------------------------------------------
            # SETTINGS
            # -------------------------------------------------

            (
                obu_map,
                owner_map,
            ) = (
                self._settings_maps()
            )


            # -------------------------------------------------
            # LER LINHAS
            # -------------------------------------------------

            parsed_rows = []


            consecutive_empty = 0


            start_row = (
                header_row + 1
            )


            worksheet_max_row = (
                worksheet.max_row
                or (
                    header_row
                    +
                    self.MAX_SCAN_ROWS
                )
            )


            end_row = min(
                worksheet_max_row,
                header_row
                +
                self.MAX_SCAN_ROWS,
            )


            for row_number, cells in enumerate(
                worksheet.iter_rows(
                    min_row=start_row,
                    max_row=end_row,
                    min_col=1,
                    max_col=73,
                ),
                start=start_row,
            ):

                if not (
                    self._row_has_source_data(
                        cells
                    )
                ):

                    consecutive_empty += 1


                    if (
                        consecutive_empty
                        >=
                        self.EMPTY_ROWS_TO_STOP
                    ):
                        break


                    continue


                consecutive_empty = 0


                parsed = (
                    self._parse_row(
                        cells,
                        row_number,
                        workbook.epoch,
                    )
                )


                self._apply_derived_values(
                    parsed,
                    obu_map,
                    owner_map,
                )


                parsed_rows.append(
                    parsed
                )

            # -------------------------------------------------
            # SEM DADOS
            # -------------------------------------------------

            if not parsed_rows:

                raise HTTPException(
                    status_code=422,
                    detail=(
                        "Nenhuma linha de dados "
                        "foi encontrada na aba "
                        "CTRL GERAL."
                    ),
                )


                        
            # -------------------------------------------------
            # REGISTROS EXISTENTES NO BANCO
            # -------------------------------------------------

            existing_by_eco = (
                self._existing_rows_by_eco(
                    [
                        row.get(
                            "eco"
                        )
                        for row
                        in parsed_rows
                    ]
                )
            )


            # -------------------------------------------------
            # AÇÃO DE CADA LINHA
            # -------------------------------------------------

            seen_signatures = {}

            for row in parsed_rows:

                eco_code = row.get(
                    "eco"
                )


                signature = self._row_signature(
                    row
                )


                first_excel_row = seen_signatures.get(
                    signature
                )


                matching_record = None
                duplicate_source = None


                if row[
                    "errors"
                ]:

                    action = "ERROR"

                elif first_excel_row is not None:

                    action = "DUPLICATE"
                    duplicate_source = "FILE"

                else:

                    seen_signatures[
                        signature
                    ] = row[
                        "excel_row"
                    ]


                    if eco_code:

                        matching_record = (
                            self._find_exact_existing(
                                row,
                                existing_by_eco.get(
                                    str(
                                        eco_code
                                    )
                                    .strip()
                                    .upper(),
                                    [],
                                ),
                            )
                        )


                    if matching_record:

                        action = "DUPLICATE"
                        duplicate_source = "DATABASE"

                    else:

                        action = "NEW"


                row[
                    "action"
                ] = action


                row[
                    "existing"
                ] = (
                    {
                        "id": str(matching_record.id),
                        "item": matching_record.item,
                    }
                    if matching_record
                    else None
                )


                row[
                    "duplicate_source"
                ] = duplicate_source
            # -------------------------------------------------
            # RESUMO
            # -------------------------------------------------

            total = len(
                parsed_rows
            )


            error_rows = sum(
                1
                for row
                in parsed_rows
                if (
                    row[
                        "action"
                    ]
                    ==
                    "ERROR"
                )
            )


            new_rows = sum(
                1
                for row
                in parsed_rows
                if (
                    row[
                        "action"
                    ]
                    ==
                    "NEW"
                )
            )


            update_rows = sum(
                1
                for row
                in parsed_rows
                if (
                    row[
                        "action"
                    ]
                    ==
                    "UPDATE"
                )
            )


            duplicate_rows = sum(
                1
                for row
                in parsed_rows
                if (
                    row[
                        "action"
                    ]
                    ==
                    "DUPLICATE"
                )
            )


            warning_rows = sum(
                1
                for row
                in parsed_rows
                if row[
                    "warnings"
                ]
            )


            valid_rows = (
                total
                -
                error_rows
            )


            preview_rows = (
                parsed_rows
                if _include_all_rows
                else parsed_rows[
                    :
                    self.MAX_PREVIEW_ROWS
                ]
            )


            return {
                "file_name":
                    filename,

                "sheet":
                    self.SHEET_NAME,

                "header_row":
                    header_row,

                "total_rows":
                    total,

                "valid_rows":
                    valid_rows,

                "new_rows":
                    new_rows,

                "update_rows":
                    update_rows,

                "duplicate_rows":
                    duplicate_rows,

                "error_rows":
                    error_rows,

                "warning_rows":
                    warning_rows,

                "can_import":
                    (
                        error_rows == 0
                        and
                        new_rows > 0
                    ),

                "preview_limit":
                    self.MAX_PREVIEW_ROWS,

                "preview_truncated":
                    (
                        total
                        >
                        self.MAX_PREVIEW_ROWS
                    ),

                "ignored_calculated_columns": [
                    "GAP",
                    "MONTH",
                    "AU",
                    "GROUP",
                    "1-Delay ECO Register",
                    "ECO registration week",
                    "ECO Origem (1)",
                    "ECO Origem (2)",
                    "ECO Origem (3)",
                    "GAP Start AZ ECO",
                    "Contar Eco emitida > 1 dias",
                    "GAP AGREEMENT",
                    "GAP AGREEMENT OTHER DEPTS",
                    "GAP TOTAL AGREEMENT OTHER DEPTS",
                    "GAP 2ST APROV R&D",
                    "GAP TOTAL 2ST APROV R&D",
                    "ECO release week",
                    "AZ Gap",
                    "Contar Eco concluída > 7 dias",
                    "Contar Eco concluída > 10 dias",
                    "Total Gap",
                    "RELEASE MONTH",
                    "RELEASE YEAR",
                ],

                "rows":
                    preview_rows,
            }


        finally:

            workbook.close()


    # =========================================================
    # IMPORTAÇÃO
    # =========================================================

    def import_file(
        self,
        file_bytes,
        filename,
    ):
        """
        Importa somente linhas classificadas como NEW.

        Erros cancelam toda a importação; duplicatas são
        ignoradas e qualquer falha no banco provoca rollback.
        """
        self._ensure_permission()

        preview = self.preview(
            file_bytes,
            filename,
            _include_all_rows=True,
        )

        if preview["error_rows"] > 0:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": (
                        "A importação foi cancelada "
                        "porque existem linhas com erro."
                    ),
                    "error_rows": preview["error_rows"],
                },
            )

        new_rows = [
            row
            for row in preview["rows"]
            if row["action"] == "NEW"
        ]

        if not new_rows:
            raise HTTPException(
                status_code=409,
                detail={
                    "message": (
                        "Nenhuma linha nova foi "
                        "encontrada para importação."
                    ),
                    "duplicate_rows": preview.get(
                        "duplicate_rows",
                        0,
                    ),
                },
            )

        imported = []

        try:
            next_item = self.ecos.next_item()
            next_position = self.ecos.next_position()

            for row in new_rows:
                data = dict(row["data"])
                derived = row["derived"]

                data["au"] = derived.get("au")
                data["group_name"] = derived.get("group")
                data["month"] = derived.get("month")

                eco = Eco(
                    **data,
                    item=next_item,
                    position=next_position,
                )

                self.db.add(eco)
                self.db.flush()

                self._record_import_history(
                    eco,
                    data,
                )

                imported.append(
                    {
                        "id": str(eco.id),
                        "item": eco.item,
                        "excel_row": row["excel_row"],
                        "eco": eco.eco,
                    }
                )

                next_item += 1
                next_position += 1

            self.db.commit()

        except HTTPException:
            self.db.rollback()
            raise

        except Exception:
            self.db.rollback()
            raise

        return {
            "file_name": filename,
            "total_rows": preview["total_rows"],
            "imported_rows": len(imported),
            "duplicate_rows": preview.get("duplicate_rows", 0),
            "warning_rows": preview["warning_rows"],
            "error_rows": 0,
            "imported": imported,
        }