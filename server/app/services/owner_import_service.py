import re
from collections import defaultdict
from io import BytesIO
from zipfile import BadZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.exc import IntegrityError

from app.models.entities import OwnerGroupMapping
from app.repository.settings_repository import SettingsRepository


class OwnerImportService:
    VALID_GROUPS = {"BOM", "DEV", "HW", "MEC"}
    MAX_FILE_BYTES = 25 * 1024 * 1024

    def __init__(self, db):
        self.db = db
        self.repository = SettingsRepository(db)

    def preview_file(self, content, filename):
        return self._analyze_file(content, filename)

    def import_file(self, content, filename):
        preview = self._analyze_file(content, filename)

        if preview["error_rows"] or preview["conflict_rows"]:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": (
                        "Corrija os erros e conflitos da planilha "
                        "antes de importar. Nenhum Owner foi inserido."
                    ),
                    "preview": preview,
                },
            )

        new_rows = [
            row for row in preview["rows"]
            if row["status"] == "NOVO"
        ]

        if not new_rows:
            return {
                "imported_rows": 0,
                "duplicate_rows": preview["duplicate_rows"],
            }

        try:
            for row in new_rows:
                self.repository.create_owner_group(
                    OwnerGroupMapping(
                        owner=row["owner"],
                        group_name=row["group"],
                    )
                )

            self.db.commit()
        except IntegrityError as error:
            self.db.rollback()
            raise HTTPException(
                status_code=409,
                detail=(
                    "As relações Owner → Group foram alteradas durante "
                    "a importação. Analise a planilha novamente."
                ),
            ) from error
        except Exception:
            self.db.rollback()
            raise

        return {
            "imported_rows": len(new_rows),
            "duplicate_rows": preview["duplicate_rows"],
        }

    def _analyze_file(self, content, filename):
        if not filename or not filename.lower().endswith(".xlsx"):
            raise HTTPException(
                status_code=400,
                detail="Formato inválido. Envie um arquivo .xlsx.",
            )

        if not content:
            raise HTTPException(
                status_code=400,
                detail="O arquivo enviado está vazio.",
            )

        if len(content) > self.MAX_FILE_BYTES:
            raise HTTPException(
                status_code=413,
                detail="O arquivo excede o limite máximo de 25 MB.",
            )

        workbook = None
        try:
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
        except (
            BadZipFile,
            InvalidFileException,
            OSError,
            ValueError,
            KeyError,
        ) as error:
            raise HTTPException(
                status_code=400,
                detail="Não foi possível ler o arquivo XLSX enviado.",
            ) from error

        try:
            if not header:
                raise HTTPException(
                    status_code=400,
                    detail="A planilha não contém as colunas OWNER e GROUP.",
                )

            header_positions = {}
            for index, value in enumerate(header):
                name = self._clean_text(value).upper()
                if name in {"OWNER", "GROUP"} and name not in header_positions:
                    header_positions[name] = index

            if not {"OWNER", "GROUP"}.issubset(header_positions):
                raise HTTPException(
                    status_code=400,
                    detail="A planilha precisa conter as colunas OWNER e GROUP.",
                )

            parsed_rows = []
            for excel_row, values in enumerate(rows, start=2):
                owner_value = self._value_at(
                    values,
                    header_positions["OWNER"],
                )
                group_value = self._value_at(
                    values,
                    header_positions["GROUP"],
                )

                if not self._clean_text(owner_value) and not self._clean_text(group_value):
                    continue

                owner = self._clean_text(owner_value)
                group = self._clean_text(group_value).upper()
                errors = []

                if not owner:
                    errors.append("Owner vazio.")
                elif len(owner) > 255:
                    errors.append("Owner excede o limite de 255 caracteres.")

                if not group:
                    errors.append("Group vazio.")
                elif group not in self.VALID_GROUPS:
                    errors.append(
                        "Group inválido. Use BOM, DEV, HW ou MEC."
                    )

                parsed_rows.append({
                    "excel_row": excel_row,
                    "owner": owner,
                    "group": group,
                    "status": "ERRO" if errors else "",
                    "message": " ".join(errors),
                    "_key": owner.casefold(),
                })
        finally:
            workbook.close()

        valid_by_owner = defaultdict(list)
        for row in parsed_rows:
            if not row["status"]:
                valid_by_owner[row["_key"]].append(row)

        existing_by_owner = defaultdict(set)
        for existing in self.repository.all_owner_group():
            existing_owner = self._clean_text(existing.owner)
            if existing_owner:
                existing_by_owner[existing_owner.casefold()].add(
                    self._clean_text(existing.group_name).upper()
                )

        for owner_key, owner_rows in valid_by_owner.items():
            file_groups = {row["group"] for row in owner_rows}

            if len(file_groups) > 1:
                for row in owner_rows:
                    row["status"] = "CONFLITO"
                    row["message"] = (
                        "O mesmo Owner aparece no arquivo com Groups diferentes."
                    )
                continue

            group = next(iter(file_groups))
            existing_groups = existing_by_owner.get(owner_key, set())

            if existing_groups:
                if existing_groups == {group}:
                    status = "DUPLICADO"
                    message = "Owner já cadastrado com o mesmo Group."
                else:
                    status = "CONFLITO"
                    message = "Owner já cadastrado com Group diferente."

                for row in owner_rows:
                    row["status"] = status
                    row["message"] = message
                continue

            owner_rows[0]["status"] = "NOVO"
            owner_rows[0]["message"] = "Novo Owner."
            for row in owner_rows[1:]:
                row["status"] = "DUPLICADO"
                row["message"] = (
                    "Owner repetido no arquivo com o mesmo Group."
                )

        counts = {
            status: sum(row["status"] == status for row in parsed_rows)
            for status in ("NOVO", "DUPLICADO", "CONFLITO", "ERRO")
        }

        for row in parsed_rows:
            row.pop("_key", None)

        return {
            "total_rows": len(parsed_rows),
            "new_rows": counts["NOVO"],
            "duplicate_rows": counts["DUPLICADO"],
            "conflict_rows": counts["CONFLITO"],
            "error_rows": counts["ERRO"],
            "can_import": (
                counts["ERRO"] == 0
                and counts["CONFLITO"] == 0
                and counts["NOVO"] > 0
            ),
            "rows": parsed_rows,
        }

    @staticmethod
    def _value_at(values, index):
        return values[index] if index < len(values) else None

    @staticmethod
    def _clean_text(value):
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value).strip())
